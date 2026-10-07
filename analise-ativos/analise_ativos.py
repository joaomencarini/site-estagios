#!/usr/bin/env python3
"""
Análise de risco e retorno de ativos brasileiros comparados ao CDI.

Como usar:
    pip install yfinance pandas numpy openpyxl requests
    python analise_ativos.py

Para analisar outros ativos, mude só a lista TICKERS logo abaixo.
O resultado é gravado em analise_ativos.xlsx (na mesma pasta do script).
"""

import sys
import time
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import requests
import yfinance as yf
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ======================================================================
# CONFIGURAÇÃO (é só aqui que você precisa mexer)
# ======================================================================
# Ticker do Yahoo Finance -> descrição (a descrição é opcional, só enfeita)
TICKERS = {
    "IVVB11.SA": "ETF S&P 500 em reais",
    "BOVA11.SA": "ETF Ibovespa",
    "WEGE3.SA": "Weg",
    "ITUB4.SA": "Itaú",
    "HGLG11.SA": "FII logístico",
}

MESES = 12                      # janela de análise
DIAS_UTEIS_ANO = 252            # para anualizar a volatilidade
ARQUIVO_SAIDA = Path(__file__).with_name("analise_ativos.xlsx")

# Ativo usado como régua nas frases da coluna "Conclusão"
# ("oscilou N vezes mais/menos que o ..."). Se ele não for baixado, usa o
# ativo de menor volatilidade. O próprio ativo de referência é comparado
# com o ativo de menor volatilidade dos demais.
ATIVO_REFERENCIA = "BOVA11.SA"

# SÓ é usado se a API do Banco Central falhar: CDI anual ESTIMADO (ajuste!).
# Quando isso acontece, o Excel e o terminal avisam que o CDI é estimativa.
CDI_ESTIMADO_ANUAL = 0.14

SERIE_CDI_SGS = 12              # série 12 do SGS: CDI diário (% ao dia)

COR_CABECALHO = "1F4E78"
CORES = ["1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD", "8C564B", "E377C2"]  # pontos do gráfico


def curto(ticker):
    """'WEGE3.SA' -> 'WEGE3' (nome curto para tabelas e frases)."""
    return ticker.replace(".SA", "")


# ======================================================================
# 1) DOWNLOAD DOS DADOS
# ======================================================================
def baixar_precos(tickers, inicio, fim):
    """Baixa o preço AJUSTADO (inclui dividendos) de cada ticker.

    Retorna (DataFrame de preços, lista de tickers que falharam).
    Se um ticker falhar, os demais continuam.
    """
    series, falhas = {}, []
    for tk in tickers:
        try:
            # auto_adjust=True: o "Close" já vem ajustado por dividendos e splits
            hist = yf.Ticker(tk).history(start=inicio, end=fim + timedelta(days=1),
                                         auto_adjust=True)
            s = hist["Close"].dropna()
            if len(s) < 30:
                raise ValueError(f"poucos dados ({len(s)} pregões)")
            s.index = pd.to_datetime(s.index).tz_localize(None).normalize()
            series[tk] = s[~s.index.duplicated()]
            print(f"  ok   {tk}: {len(s)} pregões")
        except Exception as erro:  # noqa: BLE001 - qualquer erro vira aviso
            falhas.append(tk)
            print(f"  FALHOU {tk}: {erro}")
    if not series:
        return pd.DataFrame(), falhas
    df = pd.DataFrame(series).sort_index()
    # Dias em que um ativo não negociou: repete o último preço
    # (assim todos os ativos ficam alinhados nas mesmas datas)
    df = df.ffill().dropna()
    return df, falhas


def baixar_cdi(inicio, fim):
    """Baixa o CDI diário (SGS 12) do Banco Central.

    Retorna (série de taxa diária em decimal, True se for estimativa).
    """
    url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{SERIE_CDI_SGS}/dados"
    params = {"formato": "json",
              "dataInicial": inicio.strftime("%d/%m/%Y"),
              "dataFinal": fim.strftime("%d/%m/%Y")}
    ultimo_erro = None
    for tentativa in range(1, 4):
        try:
            r = requests.get(url, params=params, timeout=30)
            r.raise_for_status()
            dados = r.json()
            if not dados:
                raise ValueError("a API devolveu lista vazia")
            s = pd.Series(
                [float(d["valor"]) / 100 for d in dados],   # % -> decimal
                index=pd.to_datetime([d["data"] for d in dados], format="%d/%m/%Y"),
            )
            print(f"  ok   CDI (SGS {SERIE_CDI_SGS}): {len(s)} dias")
            return s, False
        except Exception as erro:  # noqa: BLE001
            ultimo_erro = erro
            time.sleep(2 * tentativa)
    print(f"  FALHOU CDI (API do Banco Central): {ultimo_erro}")
    return None, True


# ======================================================================
# 2) CÁLCULOS
# ======================================================================
def calcular(precos, cdi_diario):
    """Calcula as métricas de cada ativo e do CDI (tudo em decimal)."""
    retornos = precos.pct_change().iloc[1:]
    ret_acum = precos.iloc[-1] / precos.iloc[0] - 1
    vol = retornos.std(ddof=1) * np.sqrt(DIAS_UTEIS_ANO)
    max_dd = (precos / precos.cummax() - 1).min()
    cdi_acum = (1 + cdi_diario.iloc[1:]).prod() - 1
    tabela = pd.DataFrame({
        "retorno": ret_acum,
        "volatilidade": vol,
        "acima_cdi": ret_acum - cdi_acum,
        "retorno_vol": ret_acum / vol,
        "max_drawdown": max_dd,
    })
    return tabela, cdi_acum


def frase_conclusao(tk, tabela):
    """Monta: 'O X rendeu mais/menos que o CDI, mas oscilou N vezes mais/menos que o Y'."""
    outros = [t for t in tabela.index if t != tk]
    if tk != ATIVO_REFERENCIA and ATIVO_REFERENCIA in tabela.index:
        ref = ATIVO_REFERENCIA
    elif outros:
        ref = tabela.loc[outros, "volatilidade"].idxmin()
    else:
        return f"O {curto(tk)} é o único ativo analisado."
    mais_ret = tabela.loc[tk, "acima_cdi"] >= 0
    razao = tabela.loc[tk, "volatilidade"] / tabela.loc[ref, "volatilidade"]
    mais_vol = razao >= 1
    n = razao if mais_vol else 1 / razao
    # "mas" quando há contraste (rende mais oscilando mais; rende menos oscilando menos)
    ligacao = "mas" if mais_ret == mais_vol else "e"
    return (f"O {curto(tk)} rendeu {'mais' if mais_ret else 'menos'} que o CDI, "
            f"{ligacao} oscilou {n:.1f} vezes {'mais' if mais_vol else 'menos'} "
            f"que o {curto(ref)}")


# ======================================================================
# 3) EXCEL
# ======================================================================
def estilizar_cabecalho(ws, linha, ate_coluna):
    for c in range(1, ate_coluna + 1):
        cel = ws.cell(row=linha, column=c)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def escrever_excel(precos, cdi_diario, cdi_estimado, tabela, cdi_acum, caminho):
    wb = Workbook()
    tickers = list(precos.columns)
    n_tk = len(tickers)
    n = len(precos)                      # linhas de dados
    ult = n + 1                          # última linha (cabeçalho = linha 1)
    datas = [d.to_pydatetime() for d in precos.index]
    cdi_nota = " (ESTIMADO - API do BCB falhou)" if cdi_estimado else ""

    # ---------------- Aba Preços ----------------
    ws_p = wb.active
    ws_p.title = "Preços"
    ws_p.append(["Data"] + tickers)
    for d, linha in zip(datas, precos.values):
        ws_p.append([d] + [float(v) for v in linha])
    estilizar_cabecalho(ws_p, 1, n_tk + 1)
    for r in range(2, ult + 1):
        ws_p.cell(r, 1).number_format = "dd/mm/yyyy"
        for c in range(2, n_tk + 2):
            ws_p.cell(r, c).number_format = "#,##0.00"
    ws_p.freeze_panes = "B2"
    ws_p.column_dimensions["A"].width = 13
    for c in range(2, n_tk + 2):
        ws_p.column_dimensions[get_column_letter(c)].width = 13

    # ---------------- Aba Retornos diários ----------------
    # Mesmas linhas da aba Preços: a linha 2 (1º dia) não tem retorno.
    ws_r = wb.create_sheet("Retornos diários")
    ws_r.append(["Data"] + tickers + ["CDI" + (" (estimado)" if cdi_estimado else "")])
    col_cdi = n_tk + 2
    for i, d in enumerate(datas):
        r = i + 2
        ws_r.cell(r, 1, d).number_format = "dd/mm/yyyy"
        if i == 0:
            continue
        for j in range(n_tk):
            col = get_column_letter(j + 2)
            # fórmula: preço de hoje / preço de ontem - 1
            ws_r.cell(r, j + 2, f"=Preços!{col}{r}/Preços!{col}{r-1}-1").number_format = "0.00%"
        ws_r.cell(r, col_cdi, float(cdi_diario.iloc[i])).number_format = "0.0000%"
    estilizar_cabecalho(ws_r, 1, col_cdi)
    ws_r.freeze_panes = "B2"
    ws_r.column_dimensions["A"].width = 13
    for c in range(2, col_cdi + 1):
        ws_r.column_dimensions[get_column_letter(c)].width = 14

    # ---------------- Aba Resumo ----------------
    ws = wb.create_sheet("Resumo", 0)
    cab = ["Ativo", "Descrição", "Retorno 12m", "Volatilidade anual",
           "Retorno acima do CDI", "Retorno / Volatilidade", "Maior queda (drawdown)",
           "Conclusão"]
    ws.append(cab)
    estilizar_cabecalho(ws, 1, len(cab))
    ordem = tabela.sort_values("retorno", ascending=False).index
    rng_ret = lambda col: f"'Retornos diários'!{col}3:{col}{ult}"  # noqa: E731
    linha_cdi = len(ordem) + 2
    for k, tk in enumerate(ordem):
        r = k + 2
        col = get_column_letter(tickers.index(tk) + 2)   # coluna do ativo em Preços/Retornos
        ws.cell(r, 1, tk)
        ws.cell(r, 2, TICKERS.get(tk, ""))
        ws.cell(r, 3, f"=Preços!{col}{ult}/Preços!{col}2-1")
        ws.cell(r, 4, f"=STDEV({rng_ret(col)})*SQRT({DIAS_UTEIS_ANO})")
        ws.cell(r, 5, f"=C{r}-C${linha_cdi}")
        ws.cell(r, 6, f"=C{r}/D{r}")
        ws.cell(r, 7, float(tabela.loc[tk, "max_drawdown"]))   # valor (ver nota abaixo)
        ws.cell(r, 8, frase_conclusao(tk, tabela))
    # Linha extra do CDI
    cc = get_column_letter(col_cdi)
    ws.cell(linha_cdi, 1, "CDI" + (" (ESTIMADO)" if cdi_estimado else ""))
    ws.cell(linha_cdi, 2, "Taxa DI acumulada no período" + cdi_nota)
    ws.cell(linha_cdi, 3, f"=EXP(SUMPRODUCT(LN(1+{rng_ret(cc)})))-1")
    ws.cell(linha_cdi, 4, f"=STDEV({rng_ret(cc)})*SQRT({DIAS_UTEIS_ANO})")
    ws.cell(linha_cdi, 5, "–")
    ws.cell(linha_cdi, 6, "–")
    ws.cell(linha_cdi, 7, 0.0)
    ws.cell(linha_cdi, 8, "Referência de renda fixa (sem queda de preço)")
    for c in range(1, len(cab) + 1):
        ws.cell(linha_cdi, c).font = Font(bold=True)
        ws.cell(linha_cdi, c).fill = PatternFill("solid", fgColor="DDEBF7")
    for r in range(2, linha_cdi + 1):
        for c in (3, 4, 5, 7):
            ws.cell(r, c).number_format = "0.00%"
        ws.cell(r, 6).number_format = "0.00"
        for c in (5, 6):
            ws.cell(r, c).alignment = Alignment(horizontal="right")
    # Notas
    notas = [
        f"Período: {datas[0]:%d/%m/%Y} a {datas[-1]:%d/%m/%Y} ({n} pregões). "
        "Preços ajustados (incluem dividendos e proventos).",
        "Retorno, volatilidade e retorno acima do CDI são fórmulas ligadas às abas "
        "'Preços' e 'Retornos diários'. A maior queda é um valor calculado pelo script.",
        "Volatilidade = desvio padrão dos retornos diários × raiz de 252. "
        "Retorno/Volatilidade = retorno acumulado ÷ volatilidade anual.",
    ]
    if cdi_estimado:
        notas.insert(0, f"ATENÇÃO: a API do Banco Central falhou. O CDI é uma ESTIMATIVA de "
                        f"{CDI_ESTIMADO_ANUAL:.2%} ao ano, NÃO o valor oficial.")
    for i, txt in enumerate(notas):
        c = ws.cell(linha_cdi + 2 + i, 1, txt)
        c.font = Font(italic=True, bold=cdi_estimado and i == 0,
                      color="C00000" if cdi_estimado and i == 0 else "595959")
    ws.freeze_panes = "A2"
    for letra, larg in zip("ABCDEFGH", (22, 28, 14, 14, 16, 16, 16, 88)):
        ws.column_dimensions[letra].width = larg
    ws.row_dimensions[1].height = 32

    # ---------------- Aba Gráficos ----------------
    wg = wb.create_sheet("Gráficos")
    c0 = 30                                    # dados auxiliares a partir da coluna AD
    wg.cell(1, c0, "Data")
    for j, tk in enumerate(tickers):
        wg.cell(1, c0 + 1 + j, tk)
    col_cdi_g = c0 + n_tk + 1
    wg.cell(1, col_cdi_g, "CDI acumulado" + (" (estimado)" if cdi_estimado else ""))
    for i in range(n):
        r = i + 2
        wg.cell(r, c0, f"=Preços!A{r}").number_format = "dd/mm/yyyy"
        for j in range(n_tk):
            col = get_column_letter(j + 2)
            wg.cell(r, c0 + 1 + j, f"=Preços!{col}{r}/Preços!{col}$2*100").number_format = "0.00"
        if i == 0:
            wg.cell(r, col_cdi_g, 100)
        else:
            ant = get_column_letter(col_cdi_g)
            wg.cell(r, col_cdi_g,
                    f"={ant}{r-1}*(1+'Retornos diários'!{cc}{r})").number_format = "0.00"
    for c in range(c0, col_cdi_g + 1):
        wg.cell(1, c).font = Font(bold=True)
        wg.column_dimensions[get_column_letter(c)].width = 13
    wg.cell(1, 1, "Dados dos gráficos: colunas AD em diante (base 100 = primeiro dia)")
    wg.cell(1, 1).font = Font(italic=True, color="595959")

    # (a) Linhas: preços normalizados base 100 + CDI
    g1 = LineChart()
    g1.title = "Preços normalizados (base 100) e CDI acumulado" + (" estimado" if cdi_estimado else "")
    g1.y_axis.title = "Base 100"
    g1.x_axis.number_format = "mmm/yy"
    g1.x_axis.title = "Data"
    g1.height, g1.width = 11, 26
    dados = Reference(wg, min_col=c0 + 1, max_col=col_cdi_g, min_row=1, max_row=ult)
    g1.add_data(dados, titles_from_data=True)
    g1.set_categories(Reference(wg, min_col=c0, min_row=2, max_row=ult))
    for s in g1.series:
        s.smooth = False
        s.graphicalProperties.line.width = 15000
    g1.series[-1].graphicalProperties.line.dashStyle = "dash"       # CDI tracejado
    g1.series[-1].graphicalProperties.line.solidFill = "000000"
    g1.x_axis.tickLblSkip = max(1, n // 12)
    g1.x_axis.delete = False
    g1.y_axis.delete = False
    g1.legend.position = "b"
    wg.add_chart(g1, "A3")

    # (b) Dispersão risco x retorno (1 série por ativo => legenda com o nome)
    g2 = ScatterChart()
    g2.title = "Risco x retorno (12 meses)"
    g2.style = 13
    g2.scatterStyle = "lineMarker"
    g2.x_axis.title = "Volatilidade anualizada"
    g2.y_axis.title = "Retorno acumulado"
    g2.x_axis.number_format = "0%"
    g2.y_axis.number_format = "0%"
    g2.height, g2.width = 11, 26
    for k in range(len(ordem)):
        r = k + 2
        serie = Series(Reference(ws, min_col=3, min_row=r),                 # Y = retorno
                       Reference(ws, min_col=4, min_row=r))                 # X = volatilidade
        serie.tx = SeriesLabel(strRef=None, v=ordem[k])
        serie.marker.symbol = "circle"
        serie.marker.size = 11
        cor = CORES[k % len(CORES)]
        serie.marker.graphicalProperties.solidFill = cor      # ponto preenchido
        serie.marker.graphicalProperties.line.solidFill = cor
        serie.graphicalProperties.line.noFill = True
        g2.series.append(serie)
    g2.x_axis.delete = False
    g2.y_axis.delete = False
    g2.legend.position = "r"
    wg.add_chart(g2, "A27")

    wb.calculation.fullCalcOnLoad = True    # o Excel recalcula as fórmulas ao abrir
    wb.save(caminho)


# ======================================================================
# 4) PROGRAMA PRINCIPAL
# ======================================================================
def main():
    fim = date.today()
    inicio = fim - timedelta(days=round(365.25 * MESES / 12))
    print(f"Período: {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}\n")

    print("Baixando preços ajustados (Yahoo Finance)...")
    precos, falhas = baixar_precos(list(TICKERS), inicio, fim)
    if precos.empty:
        sys.exit("Nenhum ticker foi baixado; nada a analisar.")
    if falhas:
        print(f"\nAVISO: seguindo sem os tickers que falharam: {', '.join(falhas)}")

    print("\nBaixando CDI (Banco Central, SGS 12)...")
    cdi_bruto, cdi_estimado = baixar_cdi(precos.index[0].date(), precos.index[-1].date())
    if cdi_estimado:
        print(f"AVISO: usando CDI ESTIMADO de {CDI_ESTIMADO_ANUAL:.2%} ao ano "
              "(NÃO é o valor oficial).")
        taxa_dia = (1 + CDI_ESTIMADO_ANUAL) ** (1 / DIAS_UTEIS_ANO) - 1
        cdi_diario = pd.Series(taxa_dia, index=precos.index)
    else:
        # Leva o CDI para as mesmas datas dos preços (dia sem CDI = 0)
        cdi_diario = cdi_bruto.reindex(precos.index).fillna(0.0)
        if abs(len(cdi_bruto) - len(precos)) > 5:
            print(f"AVISO: CDI tem {len(cdi_bruto)} dias e os preços {len(precos)}; confira.")

    tabela, cdi_acum = calcular(precos, cdi_diario)

    # Conferência manual: preço final / preço inicial - 1 (1º ativo)
    tk0 = tabela.index[0]
    manual = precos[tk0].iloc[-1] / precos[tk0].iloc[0] - 1
    assert abs(manual - tabela.loc[tk0, "retorno"]) < 1e-12, "retorno não confere!"
    print(f"\nConferência: {tk0}  {precos[tk0].iloc[-1]:.4f} / {precos[tk0].iloc[0]:.4f} - 1 "
          f"= {manual:.4%} (igual ao da tabela)")

    escrever_excel(precos, cdi_diario, cdi_estimado, tabela, cdi_acum, ARQUIVO_SAIDA)

    # ---- tabela no terminal ----
    mostra = tabela.sort_values("retorno", ascending=False).copy()
    mostra.loc["CDI" + (" (ESTIMADO)" if cdi_estimado else "")] = [
        cdi_acum, (cdi_diario.iloc[1:].std(ddof=1) * np.sqrt(DIAS_UTEIS_ANO)), np.nan, np.nan, 0.0]
    fmt = pd.DataFrame({
        "Retorno 12m": mostra["retorno"].map("{:.2%}".format),
        "Volatilidade": mostra["volatilidade"].map("{:.2%}".format),
        "Acima do CDI": mostra["acima_cdi"].map(lambda v: "–" if pd.isna(v) else f"{v:+.2%}"),
        "Ret/Vol": mostra["retorno_vol"].map(lambda v: "–" if pd.isna(v) else f"{v:.2f}"),
        "Max drawdown": mostra["max_drawdown"].map("{:.2%}".format),
    })
    print("\n" + fmt.to_string())
    print(f"\nArquivo gerado: {ARQUIVO_SAIDA}")

    # ---- resumo em 3 linhas ----
    melhor = tabela["retorno_vol"].idxmax()
    pior = tabela["retorno_vol"].idxmin()
    m = tabela.loc[melhor]
    print("\nRESUMO")
    print(f"1) Melhor retorno ajustado ao risco: {curto(melhor)} "
          f"(retorno/volatilidade = {m.retorno_vol:.2f}).")
    print(f"2) Rendeu {m.retorno:.2%} com volatilidade de {m.volatilidade:.2%} e queda máxima "
          f"de {m.max_drawdown:.2%}; {m.acima_cdi:+.2%} em relação ao CDI.")
    print(f"3) Pior relação retorno/risco: {curto(pior)} ({tabela.loc[pior, 'retorno_vol']:.2f}).")
    if cdi_estimado:
        print("   (CDI estimado: as comparações com o CDI são aproximadas.)")


if __name__ == "__main__":
    main()
